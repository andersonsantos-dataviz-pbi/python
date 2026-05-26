# ============================================================
# Script: Inventário de Dataflows e Status de Acesso às Datasources
# via Power BI REST API
#
# Objetivo:
#   1. Ler workspaces pela API:
#      GET https://api.powerbi.com/v1.0/myorg/groups
#
#   2. Armazenar a leitura dos workspaces em memória.
#
#   3. Para cada Workspace ID, ler os dataflows:
#      GET https://api.powerbi.com/v1.0/myorg/groups/{groupId}/dataflows
#
#   4. Armazenar a leitura dos dataflows em memória.
#
#   5. Para cada linha Workspace ID + Dataflow ID, consultar:
#      GET https://api.powerbi.com/v1.0/myorg/groups/{groupId}/dataflows/{dataflowId}/datasources
#
#   6. Criar a coluna Status Dataflow:
#      - "Normal" quando a API de datasources retornar sucesso 2xx.
#      - "Sem acesso" quando ocorrer:
#           a) 401 Unauthorized
#           b) 403 Forbidden
#           c) JSON com:
#              {
#                "error": {
#                  "code": "Unauthorized",
#                  "message": "User is not authorized"
#                }
#              }
#           d) 400 InvalidRequest: Unexpected dataflow error
#           e) Qualquer erro pontual em um dataflow específico.
#
#   7. Não interromper a execução por erro em dataflow individual.
#
#   8. Gerar tabela final com:
#      Workspace ID, Workspace Name, Dataflow ID, Dataflow Name, Status Dataflow
#
#   9. Listar o resultado no console.
#
#   10. Exportar o resultado para arquivo .xlsx.
#
# Requisitos:
#   pip install pandas requests openpyxl
# ============================================================

from __future__ import annotations

import time
import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
import requests
from requests import Response


# ============================================================
# Configurações principais
# ============================================================

TENANT_ID = "81a2db7b-b536-4905-a47b-6d892bd2f210"
CLIENT_ID = "6e486331-016d-4ac6-b2c7-9566a65a0513"
CLIENT_SECRET = "dTI8Q~8yHSt1zrYIXNIKjnUxpwO_swvU31b8VdwK"

# Escopo usado no fluxo client_credentials.
# As permissões efetivas precisam estar consentidas no App Registration.
SCOPE = "https://analysis.windows.net/powerbi/api/.default"

# Endpoint de autenticação Microsoft Entra ID
TOKEN_URL = f"https://login.microsoftonline.com/{TENANT_ID}/oauth2/v2.0/token"

# Endpoint base da Power BI REST API
POWERBI_API_BASE_URL = "https://api.powerbi.com/v1.0/myorg"

# Endpoint para listar workspaces
GROUPS_URL = f"{POWERBI_API_BASE_URL}/groups"

# Quantidade de registros por página na leitura de workspaces
PAGE_SIZE = 5000

# Quantidade máxima de tentativas para erros transitórios
MAX_RETRIES = 6

# Pausa base usada em retries
BASE_SLEEP_SECONDS = 2

# Pausa curta entre chamadas para reduzir risco de throttling
REQUEST_INTERVAL_SECONDS = 0.25

# Diretório de saída
OUTPUT_DIR = Path(r"D:\Dataside\RD Saúde")

# Arquivo Excel de saída
OUTPUT_FILE = OUTPUT_DIR / "PowerBI_Dataflows_Status_Datasources.xlsx"

# Nome da aba no Excel
OUTPUT_SHEET_NAME = "Dataflows_Status"

# Nome da tabela no Excel
OUTPUT_TABLE_NAME = "Listagem_Dataflows_Status"


# ============================================================
# Configuração de log
# ============================================================

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)


# ============================================================
# Funções auxiliares
# ============================================================

def safe_str(value: Any) -> str:
    """
    Converte valores para string de forma segura.

    Quando o valor for None, retorna string vazia.
    Também remove espaços no início e no fim.
    """
    if value is None:
        return ""
    return str(value).strip()


def get_access_token() -> str:
    """
    Obtém o token de acesso no Microsoft Entra ID usando client_credentials.

    O token retornado será usado no header Authorization das chamadas
    para a Power BI REST API.
    """
    payload = {
        "grant_type": "client_credentials",
        "client_id": CLIENT_ID,
        "client_secret": CLIENT_SECRET,
        "scope": SCOPE,
    }

    headers = {
        "Content-Type": "application/x-www-form-urlencoded"
    }

    logging.info("Solicitando token de acesso no Microsoft Entra ID...")

    response = requests.post(
        TOKEN_URL,
        data=payload,
        headers=headers,
        timeout=60,
    )

    if response.status_code != 200:
        raise RuntimeError(
            "Falha ao obter token de acesso. "
            f"Status: {response.status_code}. "
            f"Resposta: {response.text}"
        )

    token = response.json().get("access_token")

    if not token:
        raise RuntimeError(
            "Token de acesso não encontrado na resposta do Microsoft Entra ID."
        )

    logging.info("Token de acesso obtido com sucesso.")
    return token


def parse_response_json(response: Response) -> Dict[str, Any]:
    """
    Tenta converter a resposta HTTP em JSON.

    Caso a resposta não seja JSON válido, retorna um dicionário
    contendo o conteúdo textual da resposta.
    """
    try:
        payload = response.json()

        if isinstance(payload, dict):
            return payload

        return {"raw": payload}

    except Exception:
        return {"raw": response.text}


def parse_error_message(response: Response) -> str:
    """
    Extrai uma mensagem de erro legível da resposta da API.

    A Power BI REST API pode retornar erros em formatos diferentes:
      - {"error": {"code": "...", "message": "..."}}
      - {"error": "..."}
      - texto puro
    """
    payload = parse_response_json(response)

    error = payload.get("error")

    if isinstance(error, dict):
        code = safe_str(error.get("code"))
        message = safe_str(error.get("message"))

        if code and message:
            return f"{code}: {message}"
        if message:
            return message
        if code:
            return code

    if isinstance(error, str):
        return error

    message = safe_str(payload.get("message"))
    if message:
        return message

    return safe_str(payload)


def is_user_not_authorized_response(response: Response) -> bool:
    """
    Identifica exatamente o cenário:

      {
        "error": {
          "code": "Unauthorized",
          "message": "User is not authorized"
        }
      }

    Esse caso será classificado como:
      Status Dataflow = "Sem acesso"
    """
    payload = parse_response_json(response)

    error = payload.get("error")

    if not isinstance(error, dict):
        return False

    code = safe_str(error.get("code"))
    message = safe_str(error.get("message"))

    return (
        code.casefold() == "Unauthorized".casefold()
        and message.casefold() == "User is not authorized".casefold()
    )


def is_unexpected_dataflow_error_response(response: Response) -> bool:
    """
    Identifica o erro:

      Status: 400
      Mensagem: InvalidRequest: Unexpected dataflow error:

    Esse erro será classificado como:
      Status Dataflow = "Sem acesso"

    Motivo:
      Em alguns dataflows, a API de datasources pode falhar ao recuperar
      as fontes de dados, mesmo que o dataflow exista na listagem.
      Para fins de inventário, esse cenário não deve interromper o processo.
    """
    if response.status_code != 400:
        return False

    error_message = parse_error_message(response).casefold()

    return (
        "invalidrequest" in error_message
        and "unexpected dataflow error" in error_message
    )


def request_with_retry(
    method: str,
    url: str,
    headers: Dict[str, str],
    params: Optional[Dict[str, Any]] = None,
    max_retries: int = MAX_RETRIES,
    raise_on_400: bool = True,
    raise_on_401_403: bool = True,
) -> Response:
    """
    Executa uma requisição HTTP com retry e controle básico de rate limit.

    Tratamentos:
      - 2xx: sucesso.
      - 400: pode interromper ou retornar a resposta, conforme parâmetro.
      - 401/403: pode interromper ou retornar a resposta, conforme parâmetro.
      - 429: Too Many Requests, respeitando Retry-After quando existir.
      - 5xx: retry por possível falha temporária.

    Observação:
      Para a API de datasources dos dataflows, alguns erros 400, 401 e 403
      fazem parte da regra de classificação do Status Dataflow.
      Por isso, a função permite retornar a resposta sem lançar exceção.
    """
    for attempt in range(1, max_retries + 1):
        response = requests.request(
            method=method,
            url=url,
            headers=headers,
            params=params,
            timeout=120,
        )

        status_code = response.status_code

        if 200 <= status_code < 300:
            return response

        # 400 pode ser erro funcional da API de datasources.
        # Quando raise_on_400=False, a resposta é devolvida para classificação posterior.
        if status_code == 400:
            if raise_on_400:
                error_message = parse_error_message(response)
                raise RuntimeError(
                    "Erro ao chamar a Power BI REST API. "
                    f"Status: {status_code}. "
                    f"URL: {url}. "
                    f"Mensagem: {error_message}"
                )

            return response

        # 401 e 403 podem indicar ausência de permissão.
        # Quando raise_on_401_403=False, a resposta é devolvida para classificação posterior.
        if status_code in (401, 403):
            if raise_on_401_403:
                error_message = parse_error_message(response)
                raise PermissionError(
                    "Erro de permissão ao acessar a Power BI REST API. "
                    f"Status: {status_code}. "
                    f"URL: {url}. "
                    f"Mensagem: {error_message}"
                )

            return response

        # 429 indica throttling/rate limit.
        if status_code == 429:
            retry_after = response.headers.get("Retry-After")

            if retry_after and retry_after.isdigit():
                sleep_seconds = int(retry_after)
            else:
                sleep_seconds = BASE_SLEEP_SECONDS * attempt

            logging.warning(
                "Rate limit identificado pela API. "
                f"Status: 429. Tentativa {attempt}/{max_retries}. "
                f"Aguardando {sleep_seconds} segundos."
            )

            time.sleep(sleep_seconds)
            continue

        # Erros 5xx podem ser temporários.
        if 500 <= status_code <= 599:
            sleep_seconds = BASE_SLEEP_SECONDS * attempt

            logging.warning(
                "Erro temporário no serviço Power BI. "
                f"Status: {status_code}. Tentativa {attempt}/{max_retries}. "
                f"Aguardando {sleep_seconds} segundos."
            )

            time.sleep(sleep_seconds)
            continue

        error_message = parse_error_message(response)
        raise RuntimeError(
            "Erro ao chamar a Power BI REST API. "
            f"Status: {status_code}. "
            f"URL: {url}. "
            f"Mensagem: {error_message}"
        )

    raise RuntimeError(
        f"Falha após {max_retries} tentativas ao chamar a URL: {url}"
    )


# ============================================================
# Leitura de Workspaces
# ============================================================

def get_powerbi_groups(access_token: str) -> List[Dict[str, Any]]:
    """
    Lê todos os workspaces retornados pela API Groups.

    A API retorna um objeto com a propriedade 'value',
    contendo a lista de workspaces.

    Paginação:
      - $top: quantidade máxima por página.
      - $skip: quantidade de registros ignorados.
    """
    headers = {
        "Authorization": f"Bearer {access_token}",
        "Accept": "application/json",
    }

    all_groups: List[Dict[str, Any]] = []
    skip = 0

    logging.info("Iniciando leitura da API Groups...")

    while True:
        params = {
            "$top": PAGE_SIZE,
            "$skip": skip,
        }

        logging.info(
            f"Lendo workspaces com $top={PAGE_SIZE} e $skip={skip}..."
        )

        response = request_with_retry(
            method="GET",
            url=GROUPS_URL,
            headers=headers,
            params=params,
            raise_on_400=True,
            raise_on_401_403=True,
        )

        payload = response.json()
        groups = payload.get("value", [])

        if not isinstance(groups, list):
            raise RuntimeError(
                "Formato inesperado na resposta da API Groups. "
                "A propriedade 'value' não é uma lista."
            )

        all_groups.extend(groups)

        logging.info(
            f"Workspaces recebidos nesta página: {len(groups)}. "
            f"Total acumulado: {len(all_groups)}."
        )

        if len(groups) < PAGE_SIZE:
            break

        skip += PAGE_SIZE
        time.sleep(REQUEST_INTERVAL_SECONDS)

    logging.info(f"Leitura de workspaces concluída. Total bruto: {len(all_groups)}.")
    return all_groups


def transform_groups_to_dataframe(groups: List[Dict[str, Any]]) -> pd.DataFrame:
    """
    Transforma a lista de workspaces em DataFrame.

    Tratamentos:
      - Mantém somente id e name.
      - Renomeia id para Workspace ID.
      - Renomeia name para Workspace Name.
      - Remove o workspace Admin monitoring.
      - Remove registros sem Workspace ID.
    """
    final_columns = ["Workspace ID", "Workspace Name"]

    if not groups:
        return pd.DataFrame(columns=final_columns)

    df = pd.DataFrame(groups)

    for column in ["id", "name"]:
        if column not in df.columns:
            df[column] = ""

    df = df[["id", "name"]].copy()

    df.rename(
        columns={
            "id": "Workspace ID",
            "name": "Workspace Name",
        },
        inplace=True,
    )

    df["Workspace ID"] = df["Workspace ID"].apply(safe_str)
    df["Workspace Name"] = df["Workspace Name"].apply(safe_str)

    df = df[df["Workspace ID"] != ""].copy()

    df = df[
        df["Workspace Name"].str.casefold() != "admin monitoring".casefold()
    ].copy()

    df.sort_values(
        by=["Workspace Name", "Workspace ID"],
        ascending=[True, True],
        inplace=True,
    )

    df.reset_index(drop=True, inplace=True)

    return df[final_columns]


# ============================================================
# Leitura de Dataflows por Workspace
# ============================================================

def get_dataflows_by_workspace(
    access_token: str,
    workspace_id: str,
    workspace_name: str,
) -> List[Dict[str, Any]]:
    """
    Lê os dataflows de um workspace específico.

    O parâmetro groupId da API recebe o Workspace ID lido previamente
    pela API Groups.
    """
    headers = {
        "Authorization": f"Bearer {access_token}",
        "Accept": "application/json",
    }

    url = f"{POWERBI_API_BASE_URL}/groups/{workspace_id}/dataflows"

    logging.info(
        f"Lendo dataflows do workspace: {workspace_name} | {workspace_id}"
    )

    response = request_with_retry(
        method="GET",
        url=url,
        headers=headers,
        raise_on_400=True,
        raise_on_401_403=True,
    )

    payload = response.json()
    dataflows = payload.get("value", [])

    if not isinstance(dataflows, list):
        raise RuntimeError(
            "Formato inesperado na resposta da API Dataflows. "
            f"Workspace ID: {workspace_id}. "
            "A propriedade 'value' não é uma lista."
        )

    enriched_dataflows: List[Dict[str, Any]] = []

    for dataflow in dataflows:
        if not isinstance(dataflow, dict):
            continue

        enriched_dataflows.append(
            {
                "Workspace ID": workspace_id,
                "Workspace Name": workspace_name,
                "objectId": dataflow.get("objectId"),
                "name": dataflow.get("name"),
            }
        )

    logging.info(
        f"Dataflows encontrados no workspace '{workspace_name}': "
        f"{len(enriched_dataflows)}"
    )

    return enriched_dataflows


def get_all_dataflows_from_workspaces(
    access_token: str,
    df_workspaces: pd.DataFrame,
) -> List[Dict[str, Any]]:
    """
    Percorre os workspaces armazenados em memória e consulta os dataflows.

    Versão resiliente:
      - Se um workspace específico falhar na leitura dos dataflows,
        o erro é registrado em log e o script segue para o próximo workspace.
      - Isso evita que um workspace problemático interrompa todo o inventário.
    """
    all_dataflows: List[Dict[str, Any]] = []

    total_workspaces = len(df_workspaces)

    logging.info(
        f"Iniciando leitura de dataflows para {total_workspaces} workspaces..."
    )

    for index, row in df_workspaces.iterrows():
        workspace_id = safe_str(row["Workspace ID"])
        workspace_name = safe_str(row["Workspace Name"])

        logging.info(
            f"Processando workspace {index + 1}/{total_workspaces}: "
            f"{workspace_name} | {workspace_id}"
        )

        if not workspace_id:
            logging.warning(
                f"Workspace ignorado por ausência de Workspace ID. "
                f"Workspace Name: {workspace_name}"
            )
            continue

        try:
            dataflows = get_dataflows_by_workspace(
                access_token=access_token,
                workspace_id=workspace_id,
                workspace_name=workspace_name,
            )

            all_dataflows.extend(dataflows)

        except Exception as ex:
            logging.error(
                "Erro ao consultar dataflows do workspace. "
                "A execução seguirá para o próximo workspace. "
                f"Workspace Name: {workspace_name}. "
                f"Workspace ID: {workspace_id}. "
                f"Erro: {ex}"
            )

        time.sleep(REQUEST_INTERVAL_SECONDS)

    logging.info(
        f"Leitura de dataflows concluída. Total bruto: {len(all_dataflows)}."
    )

    return all_dataflows


def transform_dataflows_to_dataframe(
    dataflows: List[Dict[str, Any]],
) -> pd.DataFrame:
    """
    Transforma a lista de dataflows em DataFrame.

    Tratamentos:
      - Mantém Workspace ID e Workspace Name.
      - Renomeia objectId para Dataflow ID.
      - Renomeia name para Dataflow Name.
      - Remove registros sem Dataflow ID.
    """
    final_columns = [
        "Workspace ID",
        "Workspace Name",
        "Dataflow ID",
        "Dataflow Name",
    ]

    if not dataflows:
        return pd.DataFrame(columns=final_columns)

    df = pd.DataFrame(dataflows)

    expected_columns = [
        "Workspace ID",
        "Workspace Name",
        "objectId",
        "name",
    ]

    for column in expected_columns:
        if column not in df.columns:
            df[column] = ""

    df = df[expected_columns].copy()

    df.rename(
        columns={
            "objectId": "Dataflow ID",
            "name": "Dataflow Name",
        },
        inplace=True,
    )

    df["Workspace ID"] = df["Workspace ID"].apply(safe_str)
    df["Workspace Name"] = df["Workspace Name"].apply(safe_str)
    df["Dataflow ID"] = df["Dataflow ID"].apply(safe_str)
    df["Dataflow Name"] = df["Dataflow Name"].apply(safe_str)

    df = df[df["Dataflow ID"] != ""].copy()

    df.sort_values(
        by=["Workspace Name", "Dataflow Name", "Dataflow ID"],
        ascending=[True, True, True],
        inplace=True,
    )

    df.reset_index(drop=True, inplace=True)

    return df[final_columns]


# ============================================================
# Leitura de Datasources por Dataflow
# ============================================================

def get_dataflow_status_by_datasources_api(
    access_token: str,
    workspace_id: str,
    dataflow_id: str,
) -> str:
    """
    Consulta a API de datasources de um dataflow específico.

    Parâmetros:
      - groupId recebe Workspace ID.
      - dataflowId recebe Dataflow ID.

    Regra de negócio:
      - Se retornar sucesso 2xx:
          Status Dataflow = "Normal"

      - Se retornar JSON:
          {
            "error": {
              "code": "Unauthorized",
              "message": "User is not authorized"
            }
          }
          Status Dataflow = "Sem acesso"

      - Se retornar 401 ou 403:
          Status Dataflow = "Sem acesso"

      - Se retornar 400 com:
          InvalidRequest: Unexpected dataflow error
          Status Dataflow = "Sem acesso"

    Importante:
      Esta função não deve interromper o processamento geral por erros
      conhecidos de permissão ou inconsistência pontual de dataflow.
    """
    headers = {
        "Authorization": f"Bearer {access_token}",
        "Accept": "application/json",
    }

    url = (
        f"{POWERBI_API_BASE_URL}/groups/{workspace_id}"
        f"/dataflows/{dataflow_id}/datasources"
    )

    response = request_with_retry(
        method="GET",
        url=url,
        headers=headers,
        raise_on_400=False,
        raise_on_401_403=False,
    )

    if 200 <= response.status_code < 300:
        return "Normal"

    if is_user_not_authorized_response(response):
        return "Sem acesso"

    if response.status_code in (401, 403):
        return "Sem acesso"

    if is_unexpected_dataflow_error_response(response):
        return "Sem acesso"

    error_message = parse_error_message(response)

    raise RuntimeError(
        "Erro inesperado ao consultar datasources do dataflow. "
        f"Status: {response.status_code}. "
        f"Workspace ID: {workspace_id}. "
        f"Dataflow ID: {dataflow_id}. "
        f"Mensagem: {error_message}"
    )


def enrich_dataflows_with_status(
    access_token: str,
    df_dataflows: pd.DataFrame,
) -> pd.DataFrame:
    """
    Percorre a tabela de dataflows armazenada em memória e consulta
    a API de datasources para cada linha.

    Atribuições:
      - groupId = Workspace ID da linha.
      - dataflowId = Dataflow ID da linha.

    Cria:
      - Status Dataflow

    Versão mais resiliente:
      - Erros esperados de permissão são classificados como "Sem acesso".
      - Erro 400 InvalidRequest: Unexpected dataflow error também é
        classificado como "Sem acesso".
      - Qualquer erro inesperado em uma linha específica não interrompe
        a leitura dos demais dataflows.
      - Todo erro pontual é registrado no log.
    """
    final_columns = [
        "Workspace ID",
        "Workspace Name",
        "Dataflow ID",
        "Dataflow Name",
        "Status Dataflow",
    ]

    if df_dataflows.empty:
        return pd.DataFrame(columns=final_columns)

    df_result = df_dataflows.copy()

    statuses: List[str] = []

    total_rows = len(df_result)

    logging.info(
        f"Iniciando validação de acesso às datasources de {total_rows} dataflows..."
    )

    for index, row in df_result.iterrows():
        workspace_id = safe_str(row["Workspace ID"])
        workspace_name = safe_str(row["Workspace Name"])
        dataflow_id = safe_str(row["Dataflow ID"])
        dataflow_name = safe_str(row["Dataflow Name"])

        logging.info(
            "Consultando datasources do dataflow "
            f"{index + 1}/{total_rows}: "
            f"{workspace_name} | {dataflow_name}"
        )

        if not workspace_id or not dataflow_id:
            statuses.append("Sem acesso")

            logging.warning(
                "Linha classificada como 'Sem acesso' por ausência de "
                "Workspace ID ou Dataflow ID. "
                f"Workspace Name: {workspace_name}. "
                f"Workspace ID: {workspace_id}. "
                f"Dataflow Name: {dataflow_name}. "
                f"Dataflow ID: {dataflow_id}."
            )

            continue

        try:
            status = get_dataflow_status_by_datasources_api(
                access_token=access_token,
                workspace_id=workspace_id,
                dataflow_id=dataflow_id,
            )

        except Exception as ex:
            status = "Sem acesso"

            logging.error(
                "Erro ao consultar datasources do dataflow. "
                "Classificando como 'Sem acesso' e seguindo para o próximo. "
                f"Workspace Name: {workspace_name}. "
                f"Workspace ID: {workspace_id}. "
                f"Dataflow Name: {dataflow_name}. "
                f"Dataflow ID: {dataflow_id}. "
                f"Erro: {ex}"
            )

        statuses.append(status)

        time.sleep(REQUEST_INTERVAL_SECONDS)

    df_result["Status Dataflow"] = statuses

    df_result = df_result[final_columns].copy()

    df_result.sort_values(
        by=["Status Dataflow", "Workspace Name", "Dataflow Name"],
        ascending=[True, True, True],
        inplace=True,
    )

    df_result.reset_index(drop=True, inplace=True)

    return df_result


# ============================================================
# Exportação para Excel
# ============================================================

def write_excel_for_powerbi(df: pd.DataFrame, output_file: Path) -> None:
    """
    Gera um arquivo Excel formatado como tabela para consumo no Power BI.

    A tabela estruturada facilita a leitura no Power BI Desktop
    e melhora a governança do arquivo de saída.
    """
    output_file.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        df.to_excel(
            writer,
            sheet_name=OUTPUT_SHEET_NAME,
            index=False,
            startrow=0,
        )

        worksheet = writer.sheets[OUTPUT_SHEET_NAME]

        from openpyxl.worksheet.table import Table, TableStyleInfo

        max_row = worksheet.max_row
        max_col = worksheet.max_column

        if max_row >= 1 and max_col >= 1:
            table_ref = f"A1:E{max_row}"

            table = Table(
                displayName=OUTPUT_TABLE_NAME,
                ref=table_ref,
            )

            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )

            table.tableStyleInfo = style
            worksheet.add_table(table)

        worksheet.column_dimensions["A"].width = 42
        worksheet.column_dimensions["B"].width = 70
        worksheet.column_dimensions["C"].width = 42
        worksheet.column_dimensions["D"].width = 70
        worksheet.column_dimensions["E"].width = 22

    logging.info(f"Arquivo Excel gerado com sucesso: {output_file}")


# ============================================================
# Execução principal
# ============================================================

def main() -> None:
    """
    Função principal de execução do script.

    Fluxo:
      1. Obtém token de acesso.
      2. Lê workspaces pela API Groups.
      3. Armazena os workspaces em memória no DataFrame df_workspaces.
      4. Lê dataflows de cada workspace.
      5. Armazena os dataflows em memória no DataFrame df_dataflows.
      6. Consulta datasources de cada dataflow.
      7. Classifica Status Dataflow como "Normal" ou "Sem acesso".
      8. Exporta o resultado para Excel.
      9. Lista o resultado no console.

    Característica principal:
      - O script não para em dataflow problemático.
      - Dataflows com erro pontual são classificados como "Sem acesso".
    """
    started_at = datetime.now()

    logging.info("============================================================")
    logging.info("Iniciando leitura de Workspaces, Dataflows e Datasources")
    logging.info("============================================================")

    access_token = get_access_token()

    # ========================================================
    # Etapa 1: leitura de workspaces
    # ========================================================
    groups = get_powerbi_groups(access_token)

    # Leitura de workspaces armazenada em memória.
    df_workspaces = transform_groups_to_dataframe(groups)

    logging.info(
        f"Workspaces armazenados em memória: {len(df_workspaces)}"
    )

    # ========================================================
    # Etapa 2: leitura de dataflows por workspace
    # ========================================================
    dataflows = get_all_dataflows_from_workspaces(
        access_token=access_token,
        df_workspaces=df_workspaces,
    )

    # Leitura de dataflows armazenada em memória.
    df_dataflows = transform_dataflows_to_dataframe(dataflows)

    logging.info(
        f"Dataflows armazenados em memória: {len(df_dataflows)}"
    )

    # ========================================================
    # Etapa 3: leitura de datasources por Workspace ID + Dataflow ID
    # ========================================================
    df_final = enrich_dataflows_with_status(
        access_token=access_token,
        df_dataflows=df_dataflows,
    )

    # ========================================================
    # Etapa 4: exportação para Excel
    # ========================================================
    write_excel_for_powerbi(df_final, OUTPUT_FILE)

    # ========================================================
    # Etapa 5: listagem do resultado
    # ========================================================
    logging.info("============================================================")
    logging.info("Resultado final")
    logging.info("============================================================")

    print("\nLista de Dataflows com Status de Acesso:\n")

    if df_final.empty:
        print("Nenhum dataflow encontrado após a leitura dos workspaces.")
    else:
        print(df_final.to_string(index=False))

    finished_at = datetime.now()
    elapsed = finished_at - started_at

    logging.info("============================================================")
    logging.info("Resumo da execução")
    logging.info("============================================================")
    logging.info(f"Total de workspaces lidos: {len(df_workspaces)}")
    logging.info(f"Total de dataflows lidos: {len(df_dataflows)}")
    logging.info(f"Total de registros exportados: {len(df_final)}")
    logging.info(f"Arquivo exportado: {OUTPUT_FILE}")
    logging.info(f"Tempo total de execução: {elapsed}")
    logging.info("Processo finalizado com sucesso.")


if __name__ == "__main__":
    main()